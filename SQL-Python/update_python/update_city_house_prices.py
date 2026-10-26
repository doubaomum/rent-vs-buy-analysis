r"""Update CREA seasonally adjusted monthly prices and HPI, raw + staging.

Install: python3 -m pip install openpyxl psycopg[binary] python-dotenv truststore
Run:     python3 .\update_city_house_prices.py
Optional local source: --source "C:\path\MLS_HPI_Oct_2026.zip"
The source must be the MLS HPI package or its Seasonally Adjusted (M).xlsx,
not the news-release chart workbook. Blank numeric cells remain NULL.
Existing undated database rows are left untouched and excluded from staging.
"""

import argparse
from datetime import date, datetime
from decimal import Decimal
from html import unescape
from io import BytesIO
import os
from pathlib import Path
import re
import ssl
import sys
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen
from zipfile import ZipFile

import truststore
from openpyxl import load_workbook
from reconciliation import new_report, check_classification, check_changed_count, save_report

PAGE_URL = "https://www.crea.ca/housing-market-stats/mls-home-price-index/hpi-tool/"
WORKBOOK_NAME = "Seasonally Adjusted (M).xlsx"
# The source uses regional markets; retain the existing database city labels.
CITY_SHEETS = {
    "Canada": "AGGREGATE",
    "Vancouver": "GREATER_VANCOUVER",
    "Calgary": "CALGARY",
    "Edmonton": "EDMONTON",
    "Toronto": "GREATER_TORONTO",
    "Ottawa": "OTTAWA",
    "Montreal": "MONTREAL_CMA",
}
VALUE_COLUMNS = (
    "composite_hpi_sa", "single_family_hpi_sa", "one_storey_hpi_sa",
    "two_storey_hpi_sa", "townhouse_hpi_sa", "apartment_hpi_sa",
    "composite_benchmark_sa", "single_family_benchmark_sa",
    "one_storey_benchmark_sa", "two_storey_benchmark_sa",
    "townhouse_benchmark_sa", "apartment_benchmark_sa",
)
MAX_BYTES = 50 * 1024 * 1024


def download(url):
    request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    context = truststore.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
    with urlopen(request, timeout=45, context=context) as response:
        body = response.read(MAX_BYTES + 1)
    if len(body) > MAX_BYTES:
        raise ValueError("The source exceeds the expected download size.")
    return body


def find_package_url(page):
    matches = re.findall(
        r'''(?:href|data-href)\s*=\s*["']([^"']*MLS_HPI_[^"']+\.zip(?:\?[^"']*)?)["']''',
        unescape(page), re.IGNORECASE,
    )
    urls = list(dict.fromkeys(urljoin(PAGE_URL, match) for match in matches))
    urls = [url for url in urls if urlparse(url).hostname == "www.crea.ca"]
    if len(urls) != 1:
        raise ValueError(
            "Could not identify one current MLS HPI package on the official page. "
            "Download its ZIP and use --source PATH."
        )
    return urls[0]


def as_period(value):
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        result = date.fromisoformat(str(value).strip())
    if result.day != 1:
        raise ValueError(f"Expected a first-of-month date, received {result}.")
    return result


def as_number(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("Unexpected boolean in a numeric column.")
    result = Decimal(str(value).strip().replace(",", ""))
    if not result.is_finite() or result <= 0:
        raise ValueError("An invalid price or index was found in the workbook.")
    return result


def parse_workbook(body):
    # XLSX is also a ZIP container: distinguish the package by workbook member.
    with ZipFile(BytesIO(body)) as archive:
        if "xl/workbook.xml" in archive.namelist():
            workbook_bytes = body
        else:
            members = [n for n in archive.namelist()
                       if n.rsplit("/", 1)[-1] == WORKBOOK_NAME]
            if len(members) != 1:
                raise ValueError(f"The package does not contain {WORKBOOK_NAME}.")
            info = archive.getinfo(members[0])
            if info.file_size > MAX_BYTES:
                raise ValueError("The monthly workbook exceeds the expected size.")
            workbook_bytes = archive.read(members[0])

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    records, ranges = {}, {}
    missing_count = undated_count = 0
    try:
        for city, sheet_name in CITY_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Missing sheet {sheet_name}. Use the MLS HPI seasonally "
                    "adjusted monthly workbook, not news-release chart data."
                )
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = tuple(str(v or "").strip().lower() for v in next(rows))
            if headers != ("date",) + VALUE_COLUMNS:
                raise ValueError(f"Unexpected columns in {sheet_name}.")
            periods = set()
            for row in rows:
                if row[0] is None or not str(row[0]).strip():
                    if any(v is not None and str(v).strip() for v in row[1:]):
                        raise ValueError(f"An undated row contains values in {sheet_name}.")
                    undated_count += 1
                    continue
                period = as_period(row[0])
                if period > date.today().replace(day=1) or period in periods:
                    raise ValueError(f"Unexpected or duplicate month in {sheet_name}: {period}.")
                values = tuple(as_number(v) for v in row[1:])
                missing_count += values.count(None)
                periods.add(period)
                records[(city, period)] = values
            if not periods:
                raise ValueError(f"No dated observations in {sheet_name}.")
            ranges[city] = (min(periods), max(periods))
    finally:
        workbook.close()
    return records, ranges, missing_count, undated_count


def numeric_record(values):
    """Normalize source and RAW numbers without rounding or losing NULLs."""
    return tuple(as_number(v) for v in values)


def sync_database(records, ranges, report):
    import psycopg
    from psycopg import sql
    from dotenv import load_dotenv

    load_dotenv(Path(__file__).resolve().parent / "API.env")
    password = os.getenv("pg_password", "")
    if not password:
        raise ValueError("Missing pg_password in API.env.")
    stats = []
    with psycopg.connect(host="localhost", port=5432, dbname="rentvsbuy",
                         user="postgres", password=password, connect_timeout=10) as conn:
        with conn.cursor() as cur:
            cur.execute(sql.SQL("SELECT city, month_date, {} FROM raw.city_house_real_prices;")
                        .format(sql.SQL(", ").join(map(sql.Identifier, VALUE_COLUMNS))))
            stored, existing_ranges, old_values = {}, {}, {}
            for city, label, *values in cur.fetchall():
                if city not in CITY_SHEETS or label is None or not str(label).strip():
                    continue
                period = as_period(label)
                if (city, period) in stored:
                    raise ValueError(f"Duplicate dated raw rows: {city}, {period}.")
                stored[(city, period)] = label
                old_values[(city, period)] = numeric_record(values)
                existing_ranges.setdefault(city, []).append(period)
            for city, periods in existing_ranges.items():
                if ranges[city][0] > min(periods) or ranges[city][1] < max(periods):
                    raise ValueError(f"The source does not cover the stored history for {city}.")

            # Partial uniqueness leaves existing undated placeholder rows intact.
            cur.execute("""
                CREATE UNIQUE INDEX IF NOT EXISTS ux_crea_city_month_dated
                ON raw.city_house_real_prices (city, month_date)
                WHERE NULLIF(BTRIM(month_date), '') IS NOT NULL;
            """)
            columns = ("city", "month_date") + VALUE_COLUMNS
            query = sql.SQL("""
                INSERT INTO raw.city_house_real_prices AS target ({columns})
                VALUES ({placeholders})
                ON CONFLICT (city, month_date)
                WHERE NULLIF(BTRIM(month_date), '') IS NOT NULL
                DO UPDATE SET {assignments}
                WHERE ({old}) IS DISTINCT FROM ({new});
            """).format(
                columns=sql.SQL(", ").join(map(sql.Identifier, columns)),
                placeholders=sql.SQL(", ").join(sql.Placeholder() for _ in columns),
                assignments=sql.SQL(", ").join(sql.SQL("{} = EXCLUDED.{}").format(
                    sql.Identifier(c), sql.Identifier(c)) for c in VALUE_COLUMNS),
                old=sql.SQL(", ").join(sql.SQL("target.{}").format(sql.Identifier(c))
                                          for c in VALUE_COLUMNS),
                new=sql.SQL(", ").join(sql.SQL("EXCLUDED.{}").format(sql.Identifier(c))
                                          for c in VALUE_COLUMNS),
            )
            incoming = [(city, stored.get((city, period), period.isoformat())) + values
                        for (city, period), values in sorted(records.items())]
            inserted = updated = unchanged = 0
            for (city, period), values in sorted(records.items()):
                prior = old_values.get((city, period))
                current = numeric_record(values)
                if prior is None:
                    inserted += 1
                elif prior != current:
                    updated += 1
                else:
                    unchanged += 1
            check_classification(len(records), inserted, updated, unchanged)
            report["checks"]["classification"] = "PASS"
            report["planned_raw_changes"] = dict(inserted=inserted, updated=updated, unchanged=unchanged)
            cur.executemany(query, incoming)
            raw_changed = cur.rowcount
            check_changed_count(raw_changed, inserted, updated)
            report["checks"]["raw_change_count"] = "PASS"
            cur.execute("""
                SELECT COUNT(*), COUNT(*) FILTER (
                    WHERE NULLIF(BTRIM(month_date), '') IS NOT NULL),
                    MAX(NULLIF(BTRIM(month_date), '')::date)
                FROM raw.city_house_real_prices;
            """)
            raw_total, raw_dated, raw_latest = cur.fetchone()

            for table, source_column, suffix in (
                ("city_house_prices", "composite_benchmark_sa", "price"),
                ("city_indexed_house_prices", "composite_hpi_sa", "price_index"),
            ):
                stage_columns = tuple(f"{city.lower()}_{suffix}" for city in CITY_SHEETS)
                stage_table = sql.Identifier("stg", table)
                cur.execute(sql.SQL("CREATE UNIQUE INDEX IF NOT EXISTS {} ON {} (date_period)")
                            .format(sql.Identifier(f"ux_{table}_period"), stage_table))
                # Preserve original pivot semantics; skip only undated raw rows.
                cur.execute(sql.SQL("""
                    INSERT INTO {table} AS target (date_period, {columns})
                    SELECT BTRIM(month_date)::date, {pivot}
                    FROM raw.city_house_real_prices
                    WHERE NULLIF(BTRIM(month_date), '') IS NOT NULL
                      AND city = ANY(%s)
                    GROUP BY BTRIM(month_date)::date
                    ON CONFLICT (date_period) DO UPDATE SET {assignments}
                    WHERE ({old}) IS DISTINCT FROM ({new});
                """).format(
                    table=stage_table,
                    columns=sql.SQL(", ").join(map(sql.Identifier, stage_columns)),
                    pivot=sql.SQL(", ").join(sql.SQL("MAX(CASE WHEN city = {} THEN {} END)")
                        .format(sql.Literal(city), sql.Identifier(source_column))
                        for city in CITY_SHEETS),
                    assignments=sql.SQL(", ").join(sql.SQL("{} = EXCLUDED.{}").format(
                        sql.Identifier(c), sql.Identifier(c)) for c in stage_columns),
                    old=sql.SQL(", ").join(sql.SQL("target.{}").format(sql.Identifier(c))
                                              for c in stage_columns),
                    new=sql.SQL(", ").join(sql.SQL("EXCLUDED.{}").format(sql.Identifier(c))
                                              for c in stage_columns),
                ), (list(CITY_SHEETS),))
                changed = cur.rowcount
                cur.execute(sql.SQL("SELECT COUNT(*), MAX(date_period) FROM {}")
                            .format(stage_table))
                total, latest = cur.fetchone()
                stats.append((table, changed, total, latest))

                # Verify each RAW month/region pivot against its STAGING table.
                # Existing staging-only historical rows are intentionally untouched.
                cur.execute(sql.SQL("""
                    WITH expected AS (
                        SELECT BTRIM(month_date)::date AS period, {pivot}
                        FROM raw.city_house_real_prices
                        WHERE NULLIF(BTRIM(month_date), '') IS NOT NULL
                          AND city = ANY(%s)
                        GROUP BY BTRIM(month_date)::date
                    )
                    SELECT COUNT(*) FROM expected e
                    LEFT JOIN {table} s ON s.date_period = e.period
                    WHERE s.date_period IS NULL
                       OR ({actual}) IS DISTINCT FROM ({expected_cols});
                """).format(
                    pivot=sql.SQL(", ").join(sql.SQL("MAX(CASE WHEN city = {} THEN {} END) AS {}")
                        .format(sql.Literal(city), sql.Identifier(source_column), sql.Identifier(col))
                        for city, col in zip(CITY_SHEETS, stage_columns)),
                    table=stage_table,
                    actual=sql.SQL(", ").join(sql.SQL("s.{}").format(sql.Identifier(c)) for c in stage_columns),
                    expected_cols=sql.SQL(", ").join(sql.SQL("e.{}").format(sql.Identifier(c)) for c in stage_columns),
                ), (list(CITY_SHEETS),))
                mismatches = cur.fetchone()[0]
                if mismatches:
                    raise ValueError(f"RAW/STAGING reconciliation failed for {table}: {mismatches} mismatches")
                report["checks"][f"raw_{table}"] = "PASS"
                print(f"RAW/STAGING {table} reconciliation: PASS")

            # Compare all 12 source fields and keys against RAW after the UPSERT.
            cur.execute(sql.SQL("SELECT city, month_date, {} FROM raw.city_house_real_prices")
                        .format(sql.SQL(", ").join(map(sql.Identifier, VALUE_COLUMNS))))
            post_values = {}
            for city, label, *values in cur.fetchall():
                if city not in CITY_SHEETS or label is None or not str(label).strip():
                    continue
                key = (city, as_period(label))
                if key in post_values:
                    raise ValueError(f"Duplicate RAW city-month during verification: {key}")
                post_values[key] = numeric_record(values)
            source_mismatches = sum(post_values.get(key) != numeric_record(values)
                                    for key, values in records.items())
            if source_mismatches:
                raise ValueError(f"CREA/RAW reconciliation failed: {source_mismatches} mismatches")
            report["checks"]["source_raw"] = "PASS"
            print(f"CREA/RAW reconciliation: PASS ({len(records)} city-months verified)")
    # All three tables commit together before reporting success.
    report["transaction_committed"] = True
    report["committed_raw_changes"] = dict(report["planned_raw_changes"])
    report["status"] = "PASS"
    report["database"] = {
        "raw_changed_rows": raw_changed, "raw_total_rows": raw_total,
        "raw_dated_rows": raw_dated, "raw_latest_date": raw_latest.isoformat() if raw_latest else None,
        "staging": {table: {"changed_rows": changed, "total_rows": total,
                           "latest_date": latest.isoformat() if latest else None}
                    for table, changed, total, latest in stats},
    }
    print(f"Inserted records: {inserted}")
    print(f"Updated records: {updated}")
    print(f"Unchanged records: {unchanged}")
    print("RAW classification check: PASS")
    print("CREA city house prices update committed successfully.")
    print(f"Raw rows inserted or updated: {raw_changed}")
    print(f"Raw total rows (including existing undated rows): {raw_total}")
    print(f"Raw dated rows: {raw_dated}")
    print(f"Raw latest date: {raw_latest}")
    for table, changed, total, latest in stats:
        print(f"Staging {table} rows inserted or updated: {changed}")
        print(f"Staging {table} total rows: {total}")
        print(f"Staging {table} latest date: {latest}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Local MLS HPI ZIP or SA monthly XLSX")
    parser.add_argument("--check-only", action="store_true", help="Parse source without database access")
    args = parser.parse_args()
    if args.source:
        body = args.source.read_bytes()
    else:
        print("Finding the latest official CREA MLS HPI package...")
        url = find_package_url(download(PAGE_URL).decode("utf-8-sig"))
        print(f"Downloading {url.rsplit('/', 1)[-1]}...")
        body = download(url)
    records, ranges, missing, undated = parse_workbook(body)
    print(f"Dated city observations processed: {len(records)}")
    print(f"Source date range: {min(r[0] for r in ranges.values())} to "
          f"{max(r[1] for r in ranges.values())}")
    print(f"Missing numeric values retained as NULL: {missing}")
    print(f"Empty undated source rows skipped: {undated}")
    if args.check_only:
        print("Source validation passed; the database was not accessed.")
        return
    report = new_report(
        source="CREA", dataset="Seasonally Adjusted Monthly MLS HPI",
        checks=("classification", "raw_change_count", "source_raw",
                "raw_city_house_prices", "raw_city_indexed_house_prices"),
    )
    report["records"].update(received=len(records), valid=len(records),
                             missing_skipped=0, rejected=0)
    report["missing_numeric_cells_retained_as_null"] = missing
    report["empty_undated_source_rows_skipped"] = undated
    report["source_period_start"] = min(r[0] for r in ranges.values()).isoformat()
    report["source_period_end"] = max(r[1] for r in ranges.values()).isoformat()
    report["reconciliation_scope"] = (
        "CREA city-month source values -> RAW 12 numeric columns; "
        "RAW dated known-city values -> price and HPI pivot staging tables. "
        "Existing undated RAW rows and staging-only legacy rows remain untouched."
    )
    try:
        sync_database(records, ranges, report)
    except Exception as error:
        report["error_type"] = type(error).__name__
        raise
    finally:
        try:
            path = save_report(report, "city_house_prices",
                               Path(__file__).resolve().parent.parent / "reconciliation_reports")
            print(f"Reconciliation report: {path}")
        except OSError as error:
            print(f"Failed to save reconciliation report: {type(error).__name__}", file=sys.stderr)
            if report["status"] == "PASS":
                raise


if __name__ == "__main__":
    try:
        main()
    except HTTPError as error:
        print(f"CREA request failed. HTTP status: {error.code}")
        sys.exit(1)
    except (URLError, TimeoutError) as error:
        print(f"The download failed: {getattr(error, 'reason', error)}")
        sys.exit(1)
    except Exception as error:
        print(f"City house prices update failed: {error}")
        sys.exit(1)
